# import tkinter as tk
# from tkinter import filedialog, messagebox
import os
import io
import ffmpeg
import speech_recognition as sr
import pytesseract
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import openpyxl
import PyPDF2
from pdf2image import convert_from_path
from PIL import Image
import fitz  # PyMuPDF
import sys
import codecs
from google import genai
from flask import Flask, request, render_template, send_file, jsonify
from werkzeug.utils import secure_filename
import uuid

app = Flask(__name__)
app.config.from_object(config['default'])

# Create uploads folder if it doesn't exist
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Konsol çıktısı için UTF-8 ayarı
sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')
sys.stderr = codecs.getwriter('utf-8')(sys.stderr.buffer, 'strict')

# Windows konsol ayarları
if os.name == 'nt':
    os.system('chcp 65001 > nul')  # UTF-8 kod sayfasını ayarla

# Ortam yolları
try:
    os.environ["PATH"] += os.pathsep + r"C:\ffmpeg\bin"
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
except Exception as e:
    print(f"Uyarı: Ortam yolları ayarlanırken hata: {e}")

# Import configuration
from config import GEMINI_API_KEY, config

# Metin okuma fonksiyonları
def read_txt(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    except UnicodeDecodeError:
        try:
            with open(file_path, 'r', encoding='latin-1') as f:
                return f.read()
        except Exception as e:
            return f"[Hata] TXT dosyası okunamadı: {e}"

def read_docx(file_path):
    try:
        text = ""
        doc = Document(file_path)
        text += '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
        
        # Gömülü görselleri OCR ile oku
        for rel in doc.part._rels.values():
            if "image" in rel.target_ref:
                try:
                    image_data = rel.target_part.blob
                    image = Image.open(io.BytesIO(image_data))
                    ocr_text = pytesseract.image_to_string(image, lang='tur+eng')
                    if ocr_text.strip():
                        text += '\n' + ocr_text.strip()
                except Exception as e:
                    text += f"\n[Uyarı - Görsel OCR başarısız]: {e}\n"
        return text
    except Exception as e:
        return f"[Hata] DOCX dosyası okunamadı: {e}"

def read_xlsx(file_path):
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = ' '.join([str(cell) for cell in row if cell is not None])
                if line.strip():
                    text += line + '\n'
        return text
    except Exception as e:
        return f"[Hata] XLSX dosyası okunamadı: {e}"

def read_pdf(file_path):
    text = ""
    try:
        # Önce normal metin çıkarma dene
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + '\n'
                except Exception as e:
                    text += f"\n[Uyarı - Sayfa metni çıkarılamadı]: {e}\n"

        # Eğer yeterli metin yoksa OCR dene
        if len(text.strip()) < 100:
            try:
                images = convert_from_path(file_path)
                for img in images:
                    ocr_text = pytesseract.image_to_string(img, lang='tur+eng')
                    text += ocr_text + '\n'
            except Exception as e:
                text += f"\n[Uyarı - PDF OCR başarısız]: {e}\n"

        # Gömülü görselleri OCR ile tara
        try:
            doc = fitz.open(file_path)
            for i in range(len(doc)):
                for img in doc.get_page_images(i):
                    try:
                        base_image = doc.extract_image(img[0])
                        image = Image.open(io.BytesIO(base_image["image"]))
                        ocr_text = pytesseract.image_to_string(image, lang='tur+eng')
                        if ocr_text.strip():
                            text += '\n' + ocr_text.strip()
                    except Exception as e:
                        text += f"\n[Uyarı - Görsel OCR başarısız]: {e}\n"
        except Exception as e:
            text += f"\n[Uyarı - PDF görsel işleme başarısız]: {e}\n"

        lines = text.splitlines()
        return '\n'.join([line.strip() for line in lines if line.strip()])
    except Exception as e:
        return f"[Hata] PDF işlenirken: {e}"

def process_audio(file_path):
    try:
        recognizer = sr.Recognizer()
        base_name, ext = os.path.splitext(file_path)
        
        # Ses dosyasını WAV formatına çevir
        if ext.lower() != '.wav':
            wav_path = base_name + '.wav'
            try:
                ffmpeg.input(file_path).output(wav_path).run(quiet=True, overwrite_output=True)
                file_path = wav_path
            except Exception as e:
                return f"[Hata] Ses dönüştürme başarısız: {e}"

        # Ses tanıma işlemi
        with sr.AudioFile(file_path) as source:
            audio = recognizer.record(source)
            try:
                return recognizer.recognize_google(audio, language='tr-TR')
            except sr.UnknownValueError:
                return "[Hata] Ses tanınamadı"
            except sr.RequestError as e:
                return f"[Hata] Google Speech Recognition servisi hatası: {e}"
    except Exception as e:
        return f"[Hata] Ses işleme başarısız: {e}"

def dosya_turu_belirle(dosya_yolu):
    if not os.path.exists(dosya_yolu):
        return "Hata", "Dosya bulunamadı"

    ext = os.path.splitext(dosya_yolu)[1].lower()
    try:
        if ext in ['.txt', '.csv']:
            text = read_txt(dosya_yolu)
        elif ext == '.docx':
            text = read_docx(dosya_yolu)
        elif ext == '.xlsx':
            text = read_xlsx(dosya_yolu)
        elif ext == '.pdf':
            text = read_pdf(dosya_yolu)
        elif ext in ['.mp3', '.wav', '.ogg', '.m4a']:
            return "Ses dosyası", process_audio(dosya_yolu)
        else:
            return "Geçersiz format", f"Desteklenmeyen dosya formatı: {ext}"

        if not text:
            return "Hata", "Dosya içeriği boş"

        temiz_metin = '\n'.join([line.strip() for line in text.splitlines() if line.strip()])
        output_path = "okunan_metin.txt"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(temiz_metin)
        return "Metin dosyası", temiz_metin
    except Exception as e:
        return "Hata", f"İşlem sırasında hata oluştu: {e}"

def yapay_zeka(metin):
    try:
        client = genai.Client(api_key=GEMINI_API_KEY)
        response = client.models.generate_content(
            model="gemini-2.0-flash",
            contents=f"Aşağıdaki metni özetle ve bu metin hakkında 10 adet çoktan seçmeli soru üret:\n{metin}"
        )
        print("\n📌 Yapay Zekâ Yanıtı:\n")
        print(response.text)
        return response.text
    except Exception as e:
        hata_mesaji = f"[Hata] Yapay zekâ işleminde: {e}"
        print(hata_mesaji)
        return hata_mesaji

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400

    if file:
        # Generate unique filename
        filename = secure_filename(file.filename)
        unique_filename = f"{uuid.uuid4()}_{filename}"
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
        file.save(filepath)

        try:
            # Process the file
            dosya_turu, icerik = dosya_turu_belirle(filepath)
            
            if dosya_turu == "Hata":
                return jsonify({'error': icerik}), 400

            # Save processed content
            output_path = os.path.join(app.config['UPLOAD_FOLDER'], f"{unique_filename}_processed.txt")
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(icerik)

            # If it's a text file, also generate AI summary
            ai_output = None
            if dosya_turu == "Metin dosyası":
                ai_output = yapay_zeka(icerik)
                ai_output_path = os.path.join(app.config['UPLOAD_FOLDER'], f"{unique_filename}_ai.txt")
                with open(ai_output_path, 'w', encoding='utf-8') as f:
                    f.write(ai_output)

            # Clean up original file
            os.remove(filepath)

            return jsonify({
                'success': True,
                'file_type': dosya_turu,
                'content': icerik,
                'ai_summary': ai_output,
                'processed_file': f"{unique_filename}_processed.txt",
                'ai_file': f"{unique_filename}_ai.txt" if ai_output else None
            })

        except Exception as e:
            return jsonify({'error': str(e)}), 500

@app.route('/download/<filename>')
def download_file(filename):
    return send_file(
        os.path.join(app.config['UPLOAD_FOLDER'], filename),
        as_attachment=True
    )

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000)
